"""Match the manually vetted lipid list against the six positive-mode serum tables.

For every row of ``lipid_vetted_list.xlsx`` the positive-mode precursor m/z is
computed from the formula, and the row counts as detected on a column when any
feature of ``../column_*/samples.csv`` lies within the m/z tolerance (15 ppm by
default). Only the precursor m/z is used; RT plays no part.

Rows whose m/z hit cannot point to the listed lipid (LOW_CONFIDENCE) go to the
low-confidence file, all others to the high-confidence file. The script also
prints the checks the README quotes: parsing and mass self-checks, tolerance
sensitivity and a decoy baseline.

    python lipid_detection.py [--tol 15] [--out-dir .]

Needs numpy, pandas and openpyxl.
"""
import argparse
import collections
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
COLUMNS = ['column_12', 'column_16', 'column_20', 'column_22', 'column_41', 'column_85']

# monoisotopic atomic masses and the electron
AM = {'C': 12.0, 'H': 1.00782503207, 'N': 14.0030740048, 'O': 15.99491461956,
      'P': 30.97376163, 'S': 31.97207100, 'Na': 22.9897692809}
E = 0.00054857990946
PROTON = AM['H'] - E
ADDUCT = {  # m/z = M + shift
    '[M+H]+': PROTON,
    '[M+NH4]+': AM['N'] + 4 * AM['H'] - E,
    '[M+Na]+': AM['Na'] - E,
    '[M+H-H2O]+': PROTON - (2 * AM['H'] + AM['O']),
    '[M-H]-': -PROTON,
    '[M+HCOO]-': AM['H'] + AM['C'] + 2 * AM['O'] + E,
}

# positive-mode precursor for each class of the negative-mode sheet
POS_ION = {c: '[M+H]+' for c in (
    'PC', 'EtherPC', 'LPC', 'EtherLPC', 'SM', 'PE', 'EtherPE', 'LPE', 'EtherLPE', 'PS',
    'Cer_NS', 'Cer_NDS', 'Cer_NP', 'HexCer_NS', 'HexCer_HS', 'HexCer_HDS', 'Hex2Cer',
    'BileAcid', 'DMPE', 'MMPE')}
POS_ION.update({c: '[M+NH4]+' for c in ('PA', 'LPA', 'PG', 'LPG', 'PI', 'LPI', 'SSulfate')})

# classes whose m/z hit may not point to the listed lipid, with their summary label
LOW_CONFIDENCE = {'DMPE': 'DMPE/MMPE', 'MMPE': 'DMPE/MMPE', 'PA': 'PA/LPA', 'LPA': 'PA/LPA', 'FA': 'FA'}
# ...because the target is exactly the [M+H]+ of this PE / LPE: (class, +carbons, +double bonds);
# FA has no positive-mode ion at all
ISOBAR = {'DMPE': ('PE', 2, 0), 'MMPE': ('PE', 1, 0), 'PA': ('PE', -2, -1), 'LPA': ('LPE', -2, -1)}
ISOBAR_FORMULA = {
    'PE': lambda n, m: f'C{n + 5}H{2*n - 2*m + 10}NO8P',
    'LPE': lambda n, m: f'C{n + 5}H{2*n - 2*m + 12}NO7P',
}

# neutral formula from sum composition (n carbons, m double bonds in the chains)
SPECIES_FORMULA = {
    'FA': lambda n, m: f'C{n}H{2*n - 2*m}O2',
    'CAR': lambda n, m: f'C{n + 7}H{2*n - 2*m + 13}NO4',
    'MG': lambda n, m: f'C{n + 3}H{2*n - 2*m + 6}O4',
    'DG': lambda n, m: f'C{n + 3}H{2*n - 2*m + 4}O5',
    'PE P-': lambda n, m: f'C{n + 5}H{2*n - 2*m + 10}NO7P',
    'CE': lambda n, m: f'C{n + 27}H{2*n - 2*m + 44}O2',
    'TG': lambda n, m: f'C{n + 3}H{2*n - 2*m + 2}O6',
}


def composition(formula):
    counts = collections.Counter()
    for el, n in re.findall(r'([A-Z][a-z]?)(\d*)', formula):
        counts[el] += int(n) if n else 1
    return counts


def mass(formula):
    return sum(AM[el] * n for el, n in composition(formula).items())


def ppm(a, b):
    return (a - b) / b * 1e6


def read_negative_sheet(ws):
    """The 492-row MS-DIAL sheet (negative ESI): one positive ion per row, by class."""
    rows, checks = [], []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if not (isinstance(r[0], (int, float)) and r[2] is not None):
            continue   # header, banner, blank and m/z-range rows
        rt, mz, name, adduct, ref_mz, formula, cls = r[1], float(r[2]), r[3], r[4], float(r[9]), r[10], r[11]
        M = mass(formula)
        calc = M + ADDUCT[adduct]
        checks.append((ppm(calc, ref_mz), ppm(mz, calc)))
        ion = POS_ION[cls]
        rows.append(dict(sheet=ws.title, xlsx_row=i, name=name, cls=cls, formula=formula,
                         neg_adduct=adduct, rt_vetted=rt, ions={ion: M + ADDUCT[ion]}))
    return rows, np.array(checks)


def read_holpie_sheet(ws):
    """The 182-row manuscript sheet: 14 FA in negative ESI, then 168 in positive ESI.

    Every precursor ion the sheet lists is a target (TG and MG list two); fragment
    and neutral-loss columns are skipped.
    """
    table = [list(r) + [None] * 20 for r in ws.iter_rows(values_only=True)]
    rows, checks = [], []
    precursors, species, polarity = [], None, 'neg'
    for i, r in enumerate(table, start=1):
        if r[0] == 'POSITIVE ESI':
            polarity = 'pos'
        elif r[0] == 'Lipid species level ':
            sub = table[i]   # the row below: m/z measured / m/z theoretical / accuracy
            precursors = []
            for c, h in enumerate(r):
                m = re.match(r'Ion \d+: (\[M[^\]]*\][+-])', str(h)) if h else None
                if m:
                    theo = next(k for k in range(c, len(sub)) if sub[k] == 'm/z theoretical')
                    meas = c if sub[c] == 'm/z measured' else None
                    precursors.append((m.group(1), meas, theo))
        elif isinstance(r[3], (int, float)) and isinstance(r[5], (int, float)) and r[1]:
            species = (r[0] or species).strip()   # isomer rows leave the species blank
            cls, n, m = re.match(r'(FA|CAR|MG|DG|PE P-|CE|TG) ?(\d+):(\d+)', species).groups()
            formula = SPECIES_FORMULA[cls](int(n), int(m))
            M = mass(formula)
            ions = {}
            for adduct, meas, theo in precursors:
                calc = M + ADDUCT[adduct]
                checks.append((species, adduct, ppm(float(r[theo]), calc),
                               ppm(float(r[meas]), calc) if meas is not None else np.nan))
                ions[adduct] = calc
            row = dict(sheet=ws.title, xlsx_row=i, name=f'{species} | {r[1]}', cls=cls,
                       formula=formula, neg_adduct=None, rt_vetted=r[3], ions=ions)
            if polarity == 'neg':   # FA: [M-H]- only, no positive-mode ion to look for
                row.update(neg_adduct='[M-H]-', ions={})
            rows.append(row)
    return rows, pd.DataFrame(checks, columns=['species', 'adduct', 'theo_vs_calc', 'meas_vs_calc'])


def load_columns(root):
    """Feature m/z of each column's serum table."""
    feats = {}
    for c in COLUMNS:
        t = pd.read_csv(root / c / 'samples.csv', skiprows=4, low_memory=False)
        feats[c] = t['Average Mz'].to_numpy(float)
    return feats


def candidates(mz, ions, tol, shift_ppm=0.0):
    """Indices of the features within tol ppm of any target ion (shifted for decoys)."""
    idx = set()
    for target in ions.values():
        t = target * (1 + shift_ppm * 1e-6)
        idx.update(np.flatnonzero(np.abs(mz - t) / t * 1e6 <= tol).tolist())
    return sorted(idx)


def n_columns(lip, feats, tol, shift_ppm=0.0):
    return np.array([sum(bool(candidates(feats[c], ions, tol, shift_ppm)) for c in COLUMNS)
                     for ions in lip.ions])


def brute_force_n_columns(lip, feats, tol):
    """The same count with plain loops, as an independent check of n_columns."""
    out = []
    for ions in lip.ions:
        out.append(sum(any(abs(f - t) / t * 1e6 <= tol for t in ions.values() for f in feats[c].tolist())
                       for c in COLUMNS))
    return np.array(out)


def tally(n):
    return {k: int((np.asarray(n) >= k).sum()) for k in (3, 4, 5, 6)}


def low_confidence_reason(row):
    """Why a row's m/z hit cannot point to the listed lipid, or None (high confidence)."""
    if row.cls == 'FA':
        return 'no positive-mode ion (listed in negative ESI only); not searched'
    if row.cls not in ISOBAR:
        return None
    partner, dn, dm = ISOBAR[row.cls]
    n, m = (int(v) for v in re.match(r'\S+ (\d+):(\d+)', row['name']).groups())
    if m + dm < 0:
        return None   # saturated PA/LPA: no PE/LPE has that composition
    (ion, mz), = row.ions.items()
    iso = mass(ISOBAR_FORMULA[partner](n + dn, m + dm)) + ADDUCT['[M+H]+']
    assert abs(ppm(mz, iso)) < 1e-3, row['name']   # the same elemental composition
    return f'its {ion} ({mz:.4f}) is the same ion as the [M+H]+ of {partner} {n + dn}:{m + dm}'


def coeluting_pc(lip, max_drt=0.1):
    """For DMPE/MMPE rows: the listed PC within max_drt min whose [M-CH3]- is the same ion.

    DMPE/MMPE [M-H]- and PC [M-CH3]- are the same ion when the PC is one CH2 heavier.
    Returns {row index: PC species name}.
    """
    pcs = lip[lip.cls.isin(['PC', 'EtherPC'])]
    pc_comp = [composition(f) for f in pcs.formula]
    found = {}
    for i, r in lip[lip.cls.isin(['DMPE', 'MMPE'])].iterrows():
        pc = composition(r.formula) + collections.Counter({'C': 1, 'H': 2})
        same = pcs[[c == pc for c in pc_comp]]
        drt = np.abs(same.rt_vetted.to_numpy(float) - float(r.rt_vetted))
        if len(same) and drt.min() <= max_drt:
            found[i] = same.name.iloc[int(drt.argmin())].split('|')[0]
    return found


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('--tol', type=float, default=15.0, help='m/z tolerance in ppm (default 15)')
    ap.add_argument('--out-dir', type=Path, default=HERE)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(HERE / 'lipid_vetted_list.xlsx', read_only=True, data_only=True)
    neg_rows, neg_chk = read_negative_sheet(wb.worksheets[0])
    pos_rows, pos_chk = read_holpie_sheet(wb.worksheets[1])
    lip = pd.DataFrame(neg_rows + pos_rows)
    lip['reason'] = [low_confidence_reason(r) for _, r in lip.iterrows()]
    high = lip.reason.isna()
    lip['label'] = lip.cls.map(LOW_CONFIDENCE).where(~high)
    feats = load_columns(HERE.parent)

    print(f'rows: {len(neg_rows)} + {len(pos_rows)} = {len(lip)}; high confidence {high.sum()}, '
          f'low confidence {(~high).sum()}', lip[~high].label.value_counts().to_dict())
    print('rows per class, 182 sheet:', lip[lip.sheet == wb.worksheets[1].title].cls.value_counts().to_dict())
    iso = lip.cls.isin(list(ISOBAR))
    print(f'check, isobars: {(iso & ~high).sum()} {"/".join(ISOBAR)} rows have a target equal to the named '
          f'PE/LPE [M+H]+; no such PE/LPE for {lip[iso & high]["name"].tolist()}')
    print(f'check, 492 sheet: computed vs MS-DIAL reference m/z max |d| = {np.abs(neg_chk[:, 0]).max():.2f} ppm; '
          f'measured vs computed: 95% within {np.percentile(np.abs(neg_chk[:, 1]), 95):.2f} ppm')
    off = pos_chk[pos_chk.theo_vs_calc.abs() > 1]
    print(f'check, 182 sheet: sheet theoretical vs computed max |d| = {pos_chk.theo_vs_calc.abs().max():.2f} ppm; '
          f'> 1 ppm: {off[["species", "adduct"]].values.tolist()}')

    lip['n_columns_mz'] = n_columns(lip, feats, args.tol)
    same = (brute_force_n_columns(lip, feats, args.tol) == lip.n_columns_mz).all()
    print(f'check, brute-force recount: {"identical" if same else "DIFFERENT"}')
    pc_of = coeluting_pc(lip)
    for i, pc in pc_of.items():
        lip.at[i, 'reason'] += f'; within 0.1 min of {pc} in the negative run, likely its [M-CH3]- in-source fragment'
    print(f'check, DMPE/MMPE within 0.1 min of the PC whose [M-CH3]- is the same ion: '
          f'{len(pc_of)} of {lip.cls.isin(["DMPE", "MMPE"]).sum()}')
    for c in COLUMNS:
        lip[f'{c}_mz'] = [int(bool(candidates(feats[c], ions, args.tol))) for ions in lip.ions]
    lip['n_candidate_features_per_column'] = [
        ';'.join(str(len(candidates(feats[c], ions, args.tol))) for c in COLUMNS) for ions in lip.ions]
    lip['target_key'] = [tuple(sorted(round(v, 4) for v in ions.values())) for ions in lip.ions]

    print(f'\ndetected on >= k of 6 columns, {args.tol:g} ppm')
    print(f'  {"all":24s} rows={len(lip):4d}', tally(lip.n_columns_mz))
    print(f'  {"high confidence":24s} rows={high.sum():4d}', tally(lip[high].n_columns_mz))
    print(f'  {"low confidence":24s} rows={(~high).sum():4d}', tally(lip[~high].n_columns_mz))
    for label, sub in lip[~high].groupby('label', sort=False):
        print(f'    {label:22s} rows={len(sub):4d}', tally(sub.n_columns_mz))
    uniq = lip[high].drop_duplicates('target_key')
    print(f'  {"high, distinct targets":24s} n={len(uniq):4d}', tally(uniq.n_columns_mz))

    print('\ntolerance sensitivity and decoy baseline (targets shifted by +-50/75/100 ppm), '
          '% of high-confidence rows')
    sub = lip[high]
    for tol in (5.0, 10.0, 15.0):
        real = tally(n_columns(sub, feats, tol))
        decoy = collections.Counter()
        for s in (-100, -75, -50, 50, 75, 100):
            decoy.update(tally(n_columns(sub, feats, tol, s)))
        print(f'  {tol:4.0f} ppm  real ' + ' '.join(f'>={k}:{v / len(sub) * 100:5.1f}' for k, v in real.items())
              + '   decoy ' + ' '.join(f'>={k}:{decoy[k] / 6 / len(sub) * 100:5.1f}' for k in (3, 4, 5, 6)))

    nf = np.array([[int(v) for v in s.split(';')] for s in lip[high].n_candidate_features_per_column])
    print(f'\nhigh-confidence detections (row x column) with >= 2 candidate features: '
          f'{(nf >= 2).sum()} of {(nf > 0).sum()}')

    out = lip.copy()
    out['positive_target_ions'] = ['; '.join(f'{k} {v:.4f}' for k, v in d.items()) for d in out.ions]
    out = out.rename(columns={'cls': 'class', 'neg_adduct': 'vetted_neg_adduct', 'rt_vetted': 'vetted_RT_min'})
    cols = (['sheet', 'xlsx_row', 'name', 'class', 'formula', 'vetted_neg_adduct', 'vetted_RT_min',
             'positive_target_ions'] + [f'{c}_mz' for c in COLUMNS]
            + ['n_columns_mz', 'n_candidate_features_per_column'])

    print('\nper class: rows, >= k columns')
    for (low, cls), s in out.groupby([~high, 'class'], sort=False):
        print(f'  {"low " if low else "high"} {cls:10s} {len(s):3d}', tally(s.n_columns_mz))

    out[high][cols].to_csv(args.out_dir / 'lipid_detection_high_confidence.csv', index=False)
    out[~high][cols[:4] + ['reason'] + cols[4:]].to_csv(
        args.out_dir / 'lipid_detection_low_confidence.csv', index=False)
    print(f'\nwrote lipid_detection_high_confidence.csv ({high.sum()} rows) and '
          f'lipid_detection_low_confidence.csv ({(~high).sum()} rows)')


if __name__ == '__main__':
    main()
